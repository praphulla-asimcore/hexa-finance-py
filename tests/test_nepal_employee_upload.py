"""Tests for the Nepal employee bulk upload (app/services/nepal_employee_upload.py).

SAFETY-CRITICAL, same as tests/test_bank_details_upload.py: nepal_employee_master
is the HNPL bank file's only bank-detail source. Rows must line up with the
right employee (ledger code / PAN / HEX id / exact name), never re-key someone,
never erase on a blank, and never accept an account number Excel may have
mangled (leading zeros dropped, 9.10011E+14).

Run: python -m pytest tests/test_nepal_employee_upload.py
"""
import csv
import io
import os

import openpyxl

from app.services.bank_details_upload import UploadError
from app.services.bank_files import np_bank_lookup
from app.services.nepal_employee_upload import apply_plan, build_template, parse_upload, plan_changes
from tests.test_bank_details_upload import FakeDB

HEADER = ["Employee Name", "Ledger Code", "PAN Number", "Bank Name", "Account Number"]
ON_FILE = {"id": "u1", "employee_name": "Bhumika Shah", "employee_ledger_code": "123-109", "pan_number": "130314620",
           "bank_name": "Laxmi Sunrise Bank Ltd", "bank_code": "2201", "bank_account_number": "2601339834045001"}


def plan_for(rows, existing=(), header=HEADER):
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(header)
    w.writerows(rows)
    return plan_changes(parse_upload("u.csv", buf.getvalue().encode()),
                        FakeDB(nepal_employee_master=[dict(e) for e in existing]))


def only(plan):
    assert len(plan["items"]) == 1, plan["items"]
    return plan["items"][0]


def test_new_employee_gets_canonical_ips_bank():
    item = only(plan_for([["Anil BK", "", "", "NIMB", "09205080041317"]]))
    assert item["action"] == "NEW"
    p = item["payload"]
    assert p["entity"] == "HNPL"
    assert p["bank_code"] == "0501" and p["bank_name"] == "Nepal Investment Mega Bank Limited"
    assert p["bank_account_number"] == "09205080041317"   # leading zero kept (CSV is text)


def test_unknown_bank_is_rejected():
    assert only(plan_for([["Anil BK", "", "", "Bank of Nowhere", "123"]]))["action"] == "REJECT"


def test_matches_existing_by_ledger_code_and_blank_never_erases():
    item = only(plan_for([["", "123-109", "", "", ""]], [ON_FILE]))
    assert item["action"] == "UNCHANGED"


def test_name_match_fills_missing_ledger_code():
    on_file = {**ON_FILE, "employee_ledger_code": None}
    item = only(plan_for([["Bhumika Shah", "123-109", "", "", ""]], [on_file]))
    assert item["action"] == "UPDATE" and item["payload"] == {"employee_ledger_code": "123-109"}


def test_upload_never_rekeys_an_employee():
    item = only(plan_for([["Bhumika Shah", "999-999", "", "", ""]], [ON_FILE]))
    assert item["action"] == "REJECT" and "123-109" in item["reason"]


def test_ledger_code_pointing_at_someone_else_by_name_is_rejected():
    item = only(plan_for([["Sneha Mandal", "123-109", "", "", ""]], [ON_FILE]))
    assert item["action"] == "REJECT" and "Bhumika Shah" in item["reason"]


def test_keys_pointing_at_two_different_employees_is_rejected():
    other = {"id": "u2", "employee_name": "Anil BK", "pan_number": "555"}
    item = only(plan_for([["Bhumika Shah", "123-109", "555", "", ""]], [ON_FILE, other]))
    assert item["action"] == "REJECT"


def test_account_change_is_flagged_sensitive():
    plan = plan_for([["", "123-109", "", "Laxmi Sunrise Limited", "1111111111"]], [ON_FILE])
    assert plan["sensitive"] == 1 and only(plan)["action"] == "UPDATE"


def test_duplicate_in_file_is_rejected():
    plan = plan_for([["Anil BK", "", "", "NIMB", "1"], ["anil bk", "", "", "NIMB", "2"]])
    assert plan["counts"]["REJECT"] == 2


def test_excel_number_account_is_rejected():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(HEADER)
    ws.append(["Anil BK", "", "", "NIMB", 9205080041317])   # typed as a number: leading 0 already lost
    buf = io.BytesIO()
    wb.save(buf)
    item = only(plan_changes(parse_upload("u.xlsx", buf.getvalue()), FakeDB()))
    assert item["action"] == "REJECT" and "Text" in item["reason"]


def test_scientific_notation_account_is_rejected():
    assert only(plan_for([["Swagat Gautam", "", "", "Sanima Bank Limited", "9.10011E+14"]]))["action"] == "REJECT"


def test_resign_date_formats():
    item = only(plan_for([["Anil BK", "26-Apr-26"]], header=["Employee Name", "Resign Date"]))
    assert item["payload"]["resign_date"] == "2026-04-26"
    assert only(plan_for([["Anil BK", "soon"]], header=["Employee Name", "Resign Date"]))["action"] == "REJECT"


def test_apply_inserts_and_updates_by_id():
    db = FakeDB(nepal_employee_master=[{**ON_FILE, "employee_ledger_code": None}])
    buf = io.StringIO()
    csv.writer(buf).writerows([HEADER, ["Bhumika Shah", "123-109", "", "", ""], ["Anil BK", "", "", "NIMB", "1"]])
    plan = plan_changes(parse_upload("u.csv", buf.getvalue().encode()), db)
    assert apply_plan(plan, db) == {"NEW": 1, "UPDATE": 1}
    ops = {(w[0], w[2].get("_where")) for w in db.writes}
    assert ("update", ("id", "u1")) in ops and ("insert", None) in ops


def test_original_master_csv_round_trips_and_catches_its_corrupted_accounts():
    """db/Nepal_Employee_Master.csv: every column is recognised, and the 11
    scientific-notation account numbers are rejected rather than loaded."""
    path = os.path.join(os.path.dirname(__file__), "..", "db", "Nepal_Employee_Master.csv")
    rows = parse_upload("m.csv", open(path, "rb").read())
    assert len(rows) == 54
    plan = plan_changes(rows, FakeDB())
    assert plan["counts"]["REJECT"] == 11
    assert all("Excel" in i["reason"] for i in plan["items"] if i["action"] == "REJECT")


def test_template_has_no_data_rows_and_parses():
    wb = openpyxl.load_workbook(io.BytesIO(build_template()))
    assert not any(any(r) for r in wb.active.iter_rows(min_row=2, values_only=True))
    assert parse_upload("t.xlsx", build_template()) == []


def test_unrecognised_file_raises():
    try:
        parse_upload("x.csv", b"foo,bar\n1,2\n")
    except UploadError:
        return
    raise AssertionError("expected UploadError")


def test_master_csv_bank_spellings_resolve():
    assert np_bank_lookup("NIMB")[0] == "0501"
    assert np_bank_lookup("Laxmi Sunrise Limited")[0] == "2201"


def test_not_required_placeholders_are_treated_as_blank():
    item = only(plan_for([["Bhumika Shah", "0", "No SSF (Intern)"]], [ON_FILE],
                         header=["Employee Name", "Citizenship Number", "SSF Number"]))
    assert item["action"] == "UNCHANGED"
