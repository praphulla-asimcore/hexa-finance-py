"""Tests for the bulk bank-details upload (app/services/bank_details_upload.py)
and the bank-code rule it depends on (bank_files.resolve_bank_code).

SAFETY-CRITICAL: uploaded rows go into consultant_bank_overrides, which wins
bank-file precedence. A row must be rejected when its name doesn't match the
person the HEX id already belongs to (the Azeean/Nazarul ID-swap defect), when
its favourite code belongs to someone else, or when a HEX id repeats; a blank
cell must never erase; and an account-number change must be flagged.

Run: python -m pytest tests/test_bank_details_upload.py
"""
import csv
import io

import openpyxl

from app.services.bank_details_upload import (
    apply_plan, build_template, parse_upload, plan_changes, UploadError,
)
from app.services.bank_files import bank_code_to_name, resolve_bank_code


class _Query:
    def __init__(self, db, table, op=None, payload=None):
        self.db, self.table, self.op, self.payload = db, table, op, payload

    def select(self, *a, **k):
        return self

    def insert(self, payload):
        return _Query(self.db, self.table, "insert", payload)

    def update(self, payload):
        return _Query(self.db, self.table, "update", payload)

    def eq(self, col, val):
        self.payload = {**self.payload, "_where": (col, val)}
        return self

    def execute(self):
        if self.op:
            self.db.writes.append((self.op, self.table, self.payload))
            return type("R", (), {"data": []})()
        return type("R", (), {"data": self.db.tables.get(self.table, [])})()


class FakeDB:
    def __init__(self, **tables):
        self.tables = tables
        self.writes = []

    def from_(self, table):
        return _Query(self, table)


def csv_bytes(rows, header=("HEX Employee ID", "Consultant Name", "ID Number", "Bank Name",
                            "Account Number", "Favourite Beneficiary Code")):
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(header)
    w.writerows(rows)
    return buf.getvalue().encode()


def plan_for(rows, db=None, **kw):
    fmt, parsed = parse_upload("u.csv", csv_bytes(rows, **kw))
    return plan_changes(fmt, parsed, db or FakeDB())


def only(plan):
    assert len(plan["items"]) == 1, plan["items"]
    return plan["items"][0]


CASE = {"parsed_data": {"entities": [{"employees": [
    {"employeeId": "HEX-0284", "name": "Nooralif Bin Mohd Rahim"},
    {"employeeId": "HEX-0027", "name": "Muhammad Nazarul"},
]}]}}


# ── template rows ─────────────────────────────────────────────────────────────

def test_new_consultant_is_inserted_with_derived_bank_code():
    item = only(plan_for([["hex-0284", "Nooralif Bin Mohd Rahim", "990101-10-1234", "Maybank",
                           "1641 2345 6789", "hs372"]], FakeDB(payroll_cases=[CASE])))
    assert item["action"] == "NEW"
    p = item["payload"]
    assert p["employee_id"] == "HEX-0284"
    assert p["bank_code"] == "MBBEMYKL"
    assert p["bank_account_number"] == "164123456789"
    assert p["id_number"] == "990101101234" and p["id_type"] == "NRIC"
    assert p["favourite_beneficiary_code"] == "HS372"


def test_name_that_does_not_match_the_hex_id_is_rejected():
    """The ID-swap defect: Azeean's row carrying Nazarul's HEX id."""
    item = only(plan_for([["HEX-0027", "Azeean Norain Nadzwani", "", "Maybank", "164043718226", "HS149"]],
                         FakeDB(payroll_cases=[CASE])))
    assert item["action"] == "REJECT"
    assert "Muhammad Nazarul" in item["reason"]


def test_favourite_code_owned_by_someone_else_is_rejected():
    db = FakeDB(payroll_cases=[CASE], consultant_bank_overrides=[
        {"employee_id": "HEX-0027", "consultant_name": "Muhammad Nazarul", "favourite_beneficiary_code": "HS372"}])
    item = only(plan_for([["HEX-0284", "Nooralif Bin Mohd Rahim", "", "", "", "HS372"]], db))
    assert item["action"] == "REJECT" and "HEX-0027" in item["reason"]


def test_same_favourite_code_for_two_people_in_one_file_is_rejected():
    plan = plan_for([["HEX-0284", "Nooralif Bin Mohd Rahim", "", "", "", "HS372"],
                     ["HEX-0285", "Someone Else", "", "", "", "HS372"]])
    assert [i["action"] for i in plan["items"]] == ["REJECT", "REJECT"]


def test_duplicate_hex_id_in_file_is_rejected():
    plan = plan_for([["HEX-0284", "Nooralif Bin Mohd Rahim", "", "Maybank", "111", ""],
                     ["HEX-0284", "Nooralif Bin Mohd Rahim", "", "Maybank", "222", ""]])
    assert plan["counts"]["REJECT"] == 2


def test_bad_inputs_are_rejected():
    plan = plan_for([["0284", "A Name", "", "", "", ""],
                     ["HEX-0001", "A Name", "", "Maybank", "16-41a", ""],
                     ["HEX-0002", "A Name", "", "Bank Of Nowhere", "123", ""],
                     ["HEX-0003", "", "", "Maybank", "123", ""]])
    assert plan["counts"]["REJECT"] == 4


def test_blank_cell_never_erases_and_unchanged_row_is_reported():
    db = FakeDB(consultant_bank_overrides=[{
        "employee_id": "HEX-0284", "consultant_name": "Nooralif Bin Mohd Rahim",
        "bank_code": "MBBEMYKL", "bank_name": "Maybank", "bank_account_number": "164123456789",
        "favourite_beneficiary_code": "HS372"}])
    item = only(plan_for([["HEX-0284", "", "", "", "", "HS372"]], db))
    assert item["action"] == "UNCHANGED"


def test_adding_fav_code_to_existing_row_is_an_update_not_sensitive():
    db = FakeDB(consultant_bank_overrides=[{
        "employee_id": "HEX-0284", "consultant_name": "Nooralif Bin Mohd Rahim",
        "bank_code": "MBBEMYKL", "bank_account_number": "164123456789"}])
    plan = plan_for([["HEX-0284", "", "", "", "", "HS372"]], db)
    item = only(plan)
    assert item["action"] == "UPDATE"
    assert item["payload"] == {"favourite_beneficiary_code": "HS372"}
    assert plan["sensitive"] == 0


def test_account_number_change_is_flagged_sensitive():
    db = FakeDB(consultant_bank_overrides=[{
        "employee_id": "HEX-0284", "consultant_name": "Nooralif Bin Mohd Rahim",
        "bank_code": "MBBEMYKL", "bank_account_number": "164123456789"}])
    plan = plan_for([["HEX-0284", "", "", "Maybank", "999999999999", ""]], db)
    assert plan["sensitive"] == 1
    assert any(c[3] for c in only(plan)["changes"])


def test_name_spelling_variant_does_not_overwrite_name_on_file():
    db = FakeDB(consultant_bank_overrides=[{"employee_id": "HEX-0284", "consultant_name": "Nooralif Bin Mohd Rahim"}])
    item = only(plan_for([["HEX-0284", "NOORALIF BIN MOHD RAHIM", "", "", "", "HS372"]], db))
    assert "consultant_name" not in item["payload"]


def test_apply_writes_only_new_and_updated_rows():
    db = FakeDB(payroll_cases=[CASE], consultant_bank_overrides=[
        {"employee_id": "HEX-0027", "consultant_name": "Muhammad Nazarul"}])
    plan = plan_for([["HEX-0284", "Nooralif Bin Mohd Rahim", "", "Maybank", "164123456789", "HS372"],
                     ["HEX-0027", "", "", "", "", "HS020"],
                     ["HEX-0027", "Azeean", "", "", "", ""]], db)  # duplicate id -> both rejected
    done = apply_plan(plan, db, "tester")
    assert done == {"NEW": 1, "UPDATE": 0}
    assert [w[0] for w in db.writes] == ["insert"]
    assert db.writes[0][2]["source"] == "BULK_UPLOAD" and db.writes[0][2]["updated_by"] == "tester"


# ── Maybank tracker format ──────────────────────────────────────────────────

def tracker_xlsx(rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Beneficiary Code Tracker_Final"
    ws.append(["Record Type\xa0\xa0", "Corporate Code\xa0\xa0"])
    ws.append(["00", "MYMHEXAMATI"])
    ws.append([])
    ws.append(["Record Type", "Beneficiary / Staff / Remitter Code\xa0\xa0", "Beneficiary / Staff / Remitter Name 1\xa0\xa0",
               "Beneficiary / Staff / Remitter Name 2", "Beneficiary / Staff / Remitter Name 3", "Deactivated Flag\xa0\xa0",
               "Beneficiary New Id Number\xa0\xa0", "Beneficiary Old ID Number", "01 - Beneficiary Business Number",
               "Police/ Army ID/ Passport No", "Email Address", "Resident Indicator", "Provider",
               "Beneficiary Account No", "01 - Beneficiary Bank Code", None, "TYPE", "Name in CSI/MSA"])
    for code, name, deact, nric, acct, bank in rows:
        ws.append(["01", code, name, None, None, deact, nric, None, None, None, None, "Y", "MY", acct, bank, None, "Consultant", None])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_tracker_rows_resolve_by_ic_then_unique_csi_name():
    db = FakeDB(payroll_cases=[CASE], consultant_master=[
        {"apex_employee_id": "HEX-0100", "consultant_name": "Abu Example", "ic_number": "910112-08-5479",
         "bank_account_number": "1"}])
    fmt, rows = parse_upload("t.xlsx", tracker_xlsx([
        ("HS001", "ABU EXAMPLE", "N", "910112085479", "158220936867", "MBBEMYKL"),   # IC hit
        ("HS372", "Nooralif Bin Mohd Rahim", "N", "990101101234", "162000000001", "CIBBMYKL"),  # CSI name
        ("HS373", "Unknown Person", "N", "000000000000", "999", "MBBEMYKL"),           # nowhere
        ("HS374", "Gone Person", "Y", "", "998", "MBBEMYKL"),                           # deactivated
    ]))
    assert fmt == "tracker" and len(rows) == 4
    plan = plan_changes(fmt, rows, db)
    by_fav = {i["fav"]: i for i in plan["items"]}
    assert by_fav["HS001"]["action"] == "NEW" and by_fav["HS001"]["employee_id"] == "HEX-0100"
    assert by_fav["HS372"]["action"] == "NEW" and by_fav["HS372"]["employee_id"] == "HEX-0284"
    assert by_fav["HS372"]["payload"]["bank_code"] == "CIBBMYKL"
    assert by_fav["HS372"]["payload"]["bank_name"] == "Cimb"
    assert by_fav["HS373"]["action"] == "REJECT"
    assert by_fav["HS374"]["action"] == "REJECT"


def test_tracker_double_registration_keeps_newest_code():
    db = FakeDB(consultant_master=[{"apex_employee_id": "HEX-0073", "consultant_name": "Savla Sagar Pankaj",
                                    "ic_number": "", "bank_account_number": "164717319677"}])
    fmt, rows = parse_upload("t.xlsx", tracker_xlsx([
        ("HS094", "SAGAR SAVLA", "N", "", "164717319677", "MBBEMYKL"),
        ("HS230", "Savla Sagar Pankaj", "N", "", "164717319677", "MBBEMYKL"),
    ]))
    by_fav = {i["fav"]: i["action"] for i in plan_changes(fmt, rows, db)["items"]}
    assert by_fav == {"HS094": "REJECT", "HS230": "NEW"}


def test_unrecognised_file_raises():
    try:
        parse_upload("x.csv", b"foo,bar\n1,2\n")
    except UploadError:
        return
    raise AssertionError("expected UploadError")


def test_template_round_trips_with_no_data_rows():
    wb = openpyxl.load_workbook(io.BytesIO(build_template()))
    # No example row that could be uploaded by mistake.
    assert not any(any(r) for r in wb.active.iter_rows(min_row=2, values_only=True))
    assert parse_upload("t.xlsx", build_template()) == ("template", [])


# ── bank code rule ──────────────────────────────────────────────────────────

def test_stored_swift_code_wins_over_bank_name():
    """Tracker-loaded override has the code; an old Talenox bank name must not override it."""
    assert resolve_bank_code({"bankCode": "MBBEMYKL", "bankName": "RHB"}) == "MBBEMYKL"


def test_code_only_record_still_gets_a_code():
    assert resolve_bank_code({"bankCode": "BIMBMYKL", "bankName": ""}) == "BIMBMYKL"


def test_hexaflow_short_label_falls_back_to_bank_name():
    assert resolve_bank_code({"bankCode": "MBB", "bankName": "Maybank"}) == "MBBEMYKL"
    assert resolve_bank_code({"bankCode": "", "bankName": "Public Bank"}) == "PBBEMYKL"


def test_bank_code_to_name():
    assert bank_code_to_name("MBBEMYKL") == "Maybank"
    assert bank_code_to_name("SCBLMYKX") == "Standard Chartered"
    assert bank_code_to_name("ZZZZZZZZ") == ""


def test_tracker_code_already_on_file_resolves_to_its_owner():
    """Chia Kit Yau: IC/account sit under two HEX ids, but HS256 was confirmed as HEX-0230."""
    db = FakeDB(consultant_master=[
        {"apex_employee_id": "HEX-0155", "consultant_name": "Chia Kit Yau", "ic_number": "", "bank_account_number": "11410100336010"},
        {"apex_employee_id": "HEX-0230", "consultant_name": "Chia Kit Yau", "ic_number": "", "bank_account_number": "11410100336010"}],
        consultant_bank_overrides=[{"employee_id": "HEX-0230", "consultant_name": "Chia Kit Yau",
                                    "favourite_beneficiary_code": "HS256", "bank_code": "RHBBMYKL",
                                    "bank_account_number": "11410100336010"}])
    fmt, rows = parse_upload("t.xlsx", tracker_xlsx([("HS256", "Chia Kit Yau", "N", "", "11410100336010", "RHBBMYKL")]))
    item = only(plan_changes(fmt, rows, db))
    assert item["employee_id"] == "HEX-0230" and item["action"] != "REJECT"


def test_ic_and_account_pointing_at_different_people_is_ambiguous():
    db = FakeDB(consultant_master=[
        {"apex_employee_id": "HEX-0001", "consultant_name": "A Person", "ic_number": "900101011111", "bank_account_number": ""},
        {"apex_employee_id": "HEX-0002", "consultant_name": "B Person", "ic_number": "", "bank_account_number": "555"}])
    fmt, rows = parse_upload("t.xlsx", tracker_xlsx([("HS900", "A Person", "N", "900101011111", "555", "MBBEMYKL")]))
    assert only(plan_changes(fmt, rows, db))["action"] == "REJECT"


def test_new_row_whose_registered_name_differs_from_name_on_file_is_rejected():
    """HS252: Ikhram's IC/account sit on a Talenox record called 'Test Employee 01'."""
    db = FakeDB(consultant_master=[{"apex_employee_id": "HEX-0213", "consultant_name": "Test Employee 01",
                                    "ic_number": "780628715029", "bank_account_number": "114879001904"}])
    fmt, rows = parse_upload("t.xlsx", tracker_xlsx([("HS252", "MOHD IKHRAM MERICAN", "N", "780628715029", "114879001904", "MBBEMYKL")]))
    item = only(plan_changes(fmt, rows, db))
    assert item["action"] == "REJECT" and "Test Employee 01" in item["reason"]


def test_existing_row_whose_registered_name_differs_is_only_warned():
    """A company account already on file (HS160 'Farihan Gigabyte Enterprise')."""
    db = FakeDB(consultant_bank_overrides=[{"employee_id": "HEX-0088", "consultant_name": "Mohd Asrul Farihan Bin Md Nor",
                                            "favourite_beneficiary_code": "HS160", "bank_code": "CIBBMYKL",
                                            "bank_account_number": "8010225797"}])
    fmt, rows = parse_upload("t.xlsx", tracker_xlsx([("HS160", "Farihan Gigabyte Enterprise", "N", "", "8010225797", "CIBBMYKL")]))
    item = only(plan_changes(fmt, rows, db))
    assert item["action"] != "REJECT" and item["warnings"]
