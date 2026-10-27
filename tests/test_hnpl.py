"""HNPL (Nepal): IPS bank list, bulk upload bank file, pay-file parsing and
the accrual per "APEX Required Accounting Entries HNPL". Synthetic people
and numbers only — no DB or Zoho."""
import asyncio
import base64
import io

import openpyxl
import pytest

from app.routers import payroll_cases as pc
from app.services import bank_files as bf
from app.services import np_accounting as npa
from app.services.parser import parse_excel_buffer


# ── IPS bank list ─────────────────────────────────────────────────────────────
@pytest.mark.parametrize("written,code,branch", [
    # spellings taken from HNPL's own bulk upload files
    ("Siddhartha Bank Limited", "1801", "1"),
    ("Sanima Bank Ltd.", "4501", "1"),
    ("Nabil Bank", "0401", "99"),
    ("Standard Chartered Bank Nepal Limited", "0601", "11"),
    ("Global IME Bank Limited", "1901", "75"),
    ("Machhapuchchhre Bank Limited", "1501", "1"),
    ("Nepal Investment Mega Bank Limited", "0501", "1"),
    ("Nepal Investment Mega Bank", "0501", "1"),
    ("NMB Bank", "2501", "9"),
    ("Laxmi Sunrise Bank Ltd", "2201", "1"),
    ("Kumari Bank Ltd.", "1601", "9900"),
    ("Citizen Bank", "2001", "1"),
])
def test_np_bank_lookup_matches_bulk_file_codes(written, code, branch):
    hit = bf.np_bank_lookup(written)
    assert hit is not None and (hit[0], hit[2]) == (code, branch)
    assert bf.bank_name_to_code(written, "NP") == code


def test_np_bank_lookup_by_code_and_unknown():
    assert bf.np_bank_lookup("anything", "601")[1] == "Standard Chartered Bank Nepal Limited"
    assert bf.np_bank_lookup("Maybank") is None
    assert bf.np_bank_lookup("") is None


def test_np_banks_listed_for_dropdown():
    names = bf.BANK_NAMES_BY_COUNTRY["NP"]
    assert len(names) == len(bf.NP_BANKS) == 61
    assert "Standard Chartered Bank Nepal Limited" in names


# ── Accrual lines ─────────────────────────────────────────────────────────────
SBC = {"name": "Sita Example", "employeeId": "901", "grossSalary": 48000, "ssfContribution": 8000,
       "ssfDeduction": 12400, "netSalary": 35600, "costCentre": "SBC and Co."}
GW = {"name": "Ram Example", "employeeId": "902", "grossSalary": 300000, "ssfContribution": 20000,
      "travelAllowance": 10000, "educationAllowance": 15000, "unpaidLeave": 2000,
      "ssfDeduction": 31000, "sst": 1000, "tds": 40000, "taxAdjustment": 5000,
      "netSalary": 221000, "costCentre": "GuardWare"}
STAFF = {"name": "Hari Example", "employeeId": "Hari Example", "grossSalary": 100000,
         "ssfContribution": 15000, "advanceDeduction": 5000, "ssfDeduction": 23250,
         "tds": 3000, "netSalary": 68750}


@pytest.fixture
def unpinned(monkeypatch):
    """Resolve accounts by name (as for an org with no pinned IDs)."""
    monkeypatch.setattr(npa, "NP_ACCOUNT_IDS", {})
    monkeypatch.setattr(pc, "NP_ACCOUNT_IDS", npa.NP_ACCOUNT_IDS)


def _totals(lines):
    dr = round(sum(a for _k, dc, a, _l in lines if dc == "debit"), 2)
    cr = round(sum(a for _k, dc, a, _l in lines if dc == "credit"), 2)
    return dr, cr


def test_sbc_accrual_lines():
    lines, problem = npa.np_journal_lines(SBC, internal=False)
    assert problem is None
    assert [(k, dc, a) for k, dc, a, _l in lines] == [
        ("csi_expense", "debit", 40000), ("csi_expense", "debit", 8000),
        ("payable", "credit", 35600), ("ssf", "credit", 12400)]


def test_gw_accrual_splits_allowances_and_books_tax_adjustment_to_tds():
    lines, problem = npa.np_journal_lines(GW, internal=False)
    assert problem is None
    by_label = {l: (k, dc, a) for k, dc, a, l in lines}
    assert by_label["Salary"] == ("csi_expense", "debit", 253000)   # gross − SSF − travel − edu − unpaid
    assert by_label["Travel Allowance"] == ("csi_expense", "debit", 10000)
    assert by_label["Education Allowance"] == ("csi_expense", "debit", 15000)
    assert by_label["TDS"] == ("tds", "credit", 45000)
    assert by_label["SST"] == ("sst", "credit", 1000)
    assert _totals(lines) == (298000, 298000)


def test_internal_accrual_uses_staff_expense_and_employee_advance():
    lines, problem = npa.np_journal_lines(STAFF, internal=True)
    assert problem is None
    keys = {(k, dc) for k, dc, _a, _l in lines}
    assert ("staff_expense", "debit") in keys and ("advance", "credit") in keys
    assert ("csi_expense", "debit") not in keys
    assert _totals(lines) == (100000, 100000)


def test_cit_and_pf_credit_their_payables():
    lines, problem = npa.np_journal_lines(
        {**SBC, "cit": 500, "deductionPF": 300, "netSalary": 34800}, internal=False)
    assert problem is None and _totals(lines) == (48000, 48000)
    assert ("cit", "credit", 500, "CIT") in lines and ("pf", "credit", 300, "Deduction PF") in lines


def test_unmapped_or_unbalanced_rows_are_refused():
    _l, problem = npa.np_journal_lines({**SBC, "forexTranslation": 500, "netSalary": 35100}, internal=False)
    assert problem and "Forex" in problem
    _l, problem = npa.np_journal_lines({**SBC, "netSalary": 30000}, internal=False)
    assert problem and "differ" in problem


def test_resolve_accounts_by_name(unpinned):
    accounts = [{"id": "1", "name": "Staff Salaries and Wages"}, {"id": "2", "name": "SSF Payable"},
                {"id": "3", "name": "Cash at Bank - SCB"}, {"id": "4", "name": "SST Payable"},
                {"id": "5", "name": "SST Payable"}]
    ids, problems = npa.resolve_np_accounts(accounts, {"staff_expense", "ssf", "bank"})
    assert ids == {"staff_expense": "1", "ssf": "2", "bank": "3"} and not problems
    _ids, problems = npa.resolve_np_accounts(accounts, {"sst", "payable"})
    assert any("more than one" in p for p in problems) and any("Salary and Wages Payable" in p for p in problems)


ZOHO_ACCOUNTS = [{"id": f"A{i}", "name": n} for i, n in enumerate([
    "Consultant Salaries and Benefits", "Staff Salaries & Wages", "Salary and Wages Payable",
    "SSF Contribution Payable", "SST Payable", "TDS Payable- Staff", "Employee Advance",
    "Cash at Bank - SCB", "APC - Consultant Salaries and Benefits"])]


class _CaseDB:
    def __init__(self):
        self.updates = []
    def from_(self, _t):
        return self
    def update(self, payload):
        self.updates.append(payload)
        return self
    def eq(self, *_a):
        return self
    def execute(self):
        return type("R", (), {"data": []})()


def _zoho(monkeypatch, posted):
    async def accounts(_org):
        return ZOHO_ACCOUNTS
    async def contacts(_org):
        return {}
    async def create_contact(_org, name, _kind):
        return "C-" + name
    async def tags(_org):
        return []   # HNPL Zoho without a Customer tag: post untagged
    async def post(_org, payload):
        posted.append(payload)
        return {"journal_id": f"J{len(posted)}"}
    monkeypatch.setattr(pc, "fetch_accounts", accounts)
    monkeypatch.setattr(pc, "fetch_contacts", contacts)
    monkeypatch.setattr(pc, "create_contact", create_contact)
    monkeypatch.setattr(pc, "fetch_reporting_tags", tags)
    monkeypatch.setattr(pc, "post_journal_entry", post)


def test_hnpl_csi_accrual_posts_one_balanced_journal_per_consultant(monkeypatch, unpinned):
    posted = []
    _zoho(monkeypatch, posted)
    kase = {"id": "c-00000001", "type": "CSI", "entity": "HNPL", "period": "202609-EOM", "reference": "R",
            "parsed_data": {"entities": [{"sheetName": "SBC", "employees": [SBC]},
                                         {"sheetName": "GW", "employees": [GW]}]}}
    db = _CaseDB()
    res = asyncio.run(pc._auto_book_accruals(kase, db))
    assert res["success"] and res["posted"] == 2, res
    assert posted[0]["journal_date"] == "2026-09-30"
    for p in posted:
        dr = sum(l["amount"] for l in p["line_items"] if l["debit_or_credit"] == "debit")
        cr = sum(l["amount"] for l in p["line_items"] if l["debit_or_credit"] == "credit")
        assert round(dr, 2) == round(cr, 2)
        assert "tags" not in p["line_items"][0]
    gw_accounts = {l["account_id"] for l in posted[1]["line_items"]}
    assert gw_accounts == {"A0", "A2", "A3", "A4", "A5"}
    assert db.updates[-1]["zoho_journal_ids"] == ["J1", "J2"]


def test_hnpl_internal_payroll_accrual(monkeypatch, unpinned):
    posted = []
    _zoho(monkeypatch, posted)
    kase = {"id": "c-00000002", "type": "PAYROLL", "entity": "HNPL", "period": "202608-EOM", "reference": "R",
            "parsed_data": {"entities": [{"sheetName": "Sheet1", "employees": [STAFF]}]}}
    res = asyncio.run(pc._auto_book_accruals_payroll(kase, _CaseDB()))
    assert res["success"], res
    assert {l["account_id"] for l in posted[0]["line_items"]} == {"A1", "A2", "A3", "A5", "A6"}


def test_hnpl_accrual_stops_before_zoho_on_missing_account(monkeypatch, unpinned):
    posted = []
    _zoho(monkeypatch, posted)
    async def accounts(_org):
        return [a for a in ZOHO_ACCOUNTS if a["name"] != "SST Payable"]
    monkeypatch.setattr(pc, "fetch_accounts", accounts)
    kase = {"id": "c-00000003", "type": "CSI", "entity": "HNPL", "period": "202609-EOM", "reference": "R",
            "parsed_data": {"entities": [{"sheetName": "GW", "employees": [GW]}]}}
    res = asyncio.run(pc._auto_book_accruals(kase, _CaseDB()))
    assert not res["success"] and "SST Payable" in res["error"] and not posted


def test_pinned_hnpl_ids_post_without_reading_the_chart(monkeypatch):
    posted = []
    _zoho(monkeypatch, posted)
    async def no_chart(_org):
        raise AssertionError("chart of accounts must not be needed")
    monkeypatch.setattr(pc, "fetch_accounts", no_chart)
    kase = {"id": "c-00000004", "type": "CSI", "entity": "HNPL", "period": "202609-EOM", "reference": "R",
            "parsed_data": {"entities": [{"sheetName": "GW", "employees": [GW]}]}}
    res = asyncio.run(pc._auto_book_accruals(kase, _CaseDB()))
    assert res["success"], res
    by_label = {l["description"].split(" – ")[-1]: l["account_id"] for l in posted[0]["line_items"]}
    assert by_label == {
        "Salary": "3424587000001025336", "Travel Allowance": "3424587000001025336",
        "Education Allowance": "3424587000001025336", "SSF Contribution": "3424587000001025336",
        "Net Salary": "3424587000000365013", "SSF Deduction": "3424587000001080011",
        "SST": "3424587000000505073", "TDS": "3424587000000505065"}


def test_hnpl_has_no_revenue_accrual():
    kase = {"id": "c", "type": "CSI", "entity": "HNPL", "period": "202609-EOM"}
    assert asyncio.run(pc._auto_book_revenue_accrual(kase, None))["not_applicable"]


# ── Parsing ───────────────────────────────────────────────────────────────────
def _xlsx(sheets):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for title, rows in sheets:
        ws = wb.create_sheet(title)
        for r in rows:
            ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_internal_pay_calculation_parses_without_ein():
    head = ["Employee", "Branch", "Department", "Designation", "Basic Salary", "Laptop Allowance",
            "SSF Contribution", "Dearness Allowance", "Gross Income", "Unpaid Leave", "Advance Pay deduction",
            "Deduction PF", "CIT", "SSF Deduction", "SST", "TDS", "Total Deduction", "Net Salary"]
    row = ["Hari Example", "Hexamatics", "Ops", "Analyst", 75000, 5000, 15000, 5000, 100000, 0, 5000,
           0, 0, 23250, 0, 3000, 31250, 68750]
    ents = parse_excel_buffer(_xlsx([("Sheet1", [head, row, [None] * 18])]), "NP")
    assert len(ents) == 1 and not ents[0]["missingColumns"]
    emp = ents[0]["employees"][0]
    assert emp["employeeId"] == "Hari Example" and emp["advanceDeduction"] == 5000
    lines, problem = npa.np_journal_lines(emp, internal=True)
    assert problem is None and _totals(lines) == (100000, 100000)


def test_gw_sheet_uses_net_after_tax_adjustment():
    head = ["Employee", "EID", "EIN", "Branch", "SSF Contribution", "Travel Allowance", "Education Allowance",
            "Gross Income", "Unpaid Leave", "Tax adjustment deduction", "SSF Deduction", "SST", "TDS",
            "Total Deduction", "Net Salary", "Net Payable After tax adjustment"]
    row = ["Ram Example", 1, 902, "GuardWare", 20000, 10000, 15000, 300000, 2000, 5000, 31000, 1000,
           40000, 74000, 226000, 221000]
    emp = parse_excel_buffer(_xlsx([("GW", [head, row])]), "NP")[0]["employees"][0]
    assert emp["netSalary"] == 221000 and emp["taxAdjustment"] == 5000
    lines, problem = npa.np_journal_lines(emp, internal=False)
    assert problem is None and _totals(lines) == (298000, 298000)


# ── Bank file ─────────────────────────────────────────────────────────────────
class _NpBankDB:
    def __init__(self, master):
        self.master, self.updates, self.table = master, [], None
    def from_(self, t):
        self.table = t
        return self
    def select(self, *_a):
        return self
    def eq(self, *_a):
        return self
    def update(self, payload):
        self.updates.append(payload)
        return self
    def execute(self):
        return type("R", (), {"data": self.master if self.table == "nepal_employee_master" else []})()


MASTER = [
    {"employee_name": "Sita Example", "bank_name": "Standard Chartered Bank Nepal Limited",
     "bank_account_number": "01234567801"},
    {"employee_name": "Ram Example", "bank_name": "Nabil Bank", "bank_account_number": "0981001750000"},
    {"employee_name": "Gita Example", "bank_name": "Some Cooperative", "bank_account_number": "555"},
]


def test_hnpl_bulk_upload_file_matches_template():
    emps = [dict(SBC), dict(GW), {"name": "Gita Example", "employeeId": "903", "netSalary": 1000.0},
            {"name": "Nobody Example", "employeeId": "904", "netSalary": 500.0}]
    kase = {"id": "c1", "reference": "CSI-HNPL-202609-EOM-001", "entity": "HNPL", "type": "CSI",
            "period": "202609-EOM", "payment_date": "2026-09-30", "check_data": {"flags": []},
            "parsed_data": {"entities": [{"sheetName": "SBC", "employees": emps[:1]},
                                         {"sheetName": "GW", "employees": emps[1:]}]}}
    db = _NpBankDB(MASTER)
    asyncio.run(bf.get_bank_file_generator("NP")(kase, db, "tester"))
    cd = db.updates[-1]["check_data"]
    assert cd["bankFormat"] == "NP_BULK" and cd["crosscheck"]["ok"], cd["crosscheck"]
    assert [m["employeeId"] for m in cd["manualPayments"]] == ["903"]
    assert [m["employeeId"] for m in cd["excludedNoBank"]] == ["904"]

    ws = openpyxl.load_workbook(io.BytesIO(base64.b64decode(cd["bankTxt"]["data"]))).worksheets[0]
    assert ws.title == "Sep" and all(ws.cell(1, c).value is None for c in range(1, 9))
    rows = [[ws.cell(r, c).value for c in range(1, 9)] for r in range(2, ws.max_row + 1)]
    assert rows == [
        ["30/09/2026", "Sita Example", "Standard Chartered Bank Nepal Limited", "0601", "11",
         "01234567801", 35600, "SBC_Sep'26"],
        ["30/09/2026", "Ram Example", "Nabil Bank Ltd.", "0401", "99", "0981001750000", 221000, "GW_Sep'26"],
    ]


def test_hnpl_internal_payroll_bank_file_remarks():
    kase = {"id": "c2", "reference": "PAYROLL-HNPL-202608-EOM-001", "entity": "HNPL", "type": "PAYROLL",
            "period": "202608-EOM", "payment_date": "2026-08-31", "check_data": {},
            "parsed_data": {"entities": [{"sheetName": "Sheet1", "employees": [
                {"name": "Sita Example", "employeeId": "Sita Example", "netSalary": 100.0}]}]}}
    db = _NpBankDB(MASTER)
    asyncio.run(bf.get_bank_file_generator("NP")(kase, db, "tester"))
    ws = openpyxl.load_workbook(io.BytesIO(base64.b64decode(
        db.updates[-1]["check_data"]["bankTxt"]["data"]))).worksheets[0]
    assert ws.title == "Aug" and ws.cell(2, 8).value == "HNPL_Aug'26"
