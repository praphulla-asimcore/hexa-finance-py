"""Bulk upload into nepal_employee_master (/admin/nepal-employees/upload) —
the Nepal (HNPL) counterpart of app/services/bank_details_upload.py, with the
same preview-then-confirm flow and the same plan shape so one page template
serves both.

Uploads only ever ADD employees. Nepal has no IC number to corroborate on, so
a row counts as already on file when its RigoHR ledger code, PAN, HEX id or
exact name matches an existing employee (the keys bank_files.match_nepal_employee
pays by); it is then skipped as "Already in database", never updated.

Account numbers typed into Excel as numbers are rejected: Excel drops leading
zeros (Anil BK's 09205080041317) and turns 15+ digit numbers into 9.10011E+14,
which is exactly how 11 accounts in the original master got corrupted.
"""
import io
import re
from collections import defaultdict
from datetime import date, datetime, timezone

import openpyxl
from openpyxl.styles import Font, PatternFill

from app.services.bank_details_upload import UploadError, _header_key, _rows_from_file, _s
from app.services.bank_files import _norm_name, np_bank_lookup

SOURCE = "BULK_UPLOAD"

FIELD_LABELS = {
    "employee_name": "Name",
    "employee_ledger_code": "Ledger code",
    "apex_employee_id": "HEX ID",
    "citizenship_number": "Citizenship no.",
    "pan_number": "PAN",
    "nationality": "Nationality",
    "bank_name": "Bank",
    "bank_code": "Bank code",
    "bank_account_name": "Account name",
    "bank_account_number": "Account no.",
    "ssf_number": "SSF no.",
    "cit_number": "CIT no.",
    "resign_date": "Resign date",
}
# Identity keys: any of these (or the exact name) already on file means "Already in database".
_KEYS = ("employee_ledger_code", "pan_number", "apex_employee_id")

TEMPLATE_HEADERS = [
    ("employee_name", "Employee Name"),
    ("employee_ledger_code", "Ledger Code"),
    ("apex_employee_id", "HEX Employee ID"),
    ("citizenship_number", "Citizenship Number"),
    ("pan_number", "PAN Number"),
    ("nationality", "Nationality"),
    ("bank_name", "Bank Name"),
    ("bank_account_name", "Account Name"),
    ("bank_account_number", "Account Number"),
    ("ssf_number", "SSF Number"),
    ("cit_number", "CIT Number"),
    ("resign_date", "Resign Date"),
]
_HEADER_ALIASES = {_header_key(label): f for f, label in TEMPLATE_HEADERS} | {
    "name": "employee_name",
    "employee ledger code": "employee_ledger_code", "ein": "employee_ledger_code",
    "apex employee id": "apex_employee_id", "employee id": "apex_employee_id",
    "pan": "pan_number", "bank": "bank_name", "bank code": "bank_code",
    "bank account number": "bank_account_number", "account no": "bank_account_number",
    "bank account name": "bank_account_name",
}
_HEX_ID_RE = re.compile(r"^HEX-\d+$")


def _parse_date(raw):
    if isinstance(raw, datetime):
        return raw.date()
    if isinstance(raw, date):
        return raw
    for fmt in ("%Y-%m-%d", "%d-%b-%y", "%d-%b-%Y", "%d/%m/%Y", "%d-%m-%Y"):
        try:
            return datetime.strptime(_s(raw), fmt).date()
        except ValueError:
            continue
    return None


def parse_upload(filename: str, content: bytes) -> list[dict]:
    for sheet in _rows_from_file(filename, content):
        for i, header in enumerate(sheet[:15]):
            fields = {_HEADER_ALIASES[k]: j for j, k in enumerate(_header_key(h) for h in header)
                      if k in _HEADER_ALIASES}
            if "employee_name" not in fields or len(fields) < 2:
                continue
            rows = []
            for n, r in enumerate(sheet[i + 1:], start=i + 2):
                raw = {f: (r[j] if j < len(r) else None) for f, j in fields.items()}
                row = {f: _s(v) for f, v in raw.items()}
                if not any(row.values()):
                    continue
                acct = raw.get("bank_account_number")
                row["row_no"] = n
                row["account_was_number"] = isinstance(acct, (int, float)) and not isinstance(acct, bool)
                resign = _parse_date(raw.get("resign_date"))   # Excel date cells arrive as datetimes
                if resign:
                    row["resign_date"] = resign.isoformat()
                rows.append(row)
            return rows
    raise UploadError("No recognisable header row. Use the Nepal template (it has an 'Employee Name' "
                      "column) or the same columns as the Nepal Employee Master.")


def _normalise(row: dict) -> tuple[dict, str]:
    v = {f: _s(row.get(f)) for f in FIELD_LABELS}
    v["apex_employee_id"] = v["apex_employee_id"].upper().replace(" ", "")
    v["bank_account_number"] = re.sub(r"[\s-]", "", v["bank_account_number"])
    # "No SSF (Intern)", "0" etc. mean "not required" — the 2026-07-31 load
    # stored those as blank, so treat them as blank (= leave as is).
    for f in ("citizenship_number", "pan_number", "ssf_number", "cit_number"):
        if v[f] in ("0", "-", "NA", "N/A") or v[f].lower().startswith("no "):
            v[f] = ""

    if row.get("account_was_number") and v["bank_account_number"]:
        return v, ("Account Number is stored as a number in Excel, so leading zeros or digits may be lost — "
                   "format the column as Text and type it again")
    if "E+" in v["bank_account_number"].upper():
        return v, f"Account Number '{v['bank_account_number']}' was corrupted by Excel — type it again"
    if v["bank_account_number"] and not v["bank_account_number"].isdigit():
        return v, "Account Number must be digits only"
    if v["apex_employee_id"] and not _HEX_ID_RE.match(v["apex_employee_id"]):
        return v, "HEX Employee ID must look like HEX-0123"
    if v["resign_date"] and not _parse_date(v["resign_date"]):
        return v, f"Resign Date '{v['resign_date']}' isn't a date (use e.g. 26-Apr-26 or 2026-04-26)"
    if v["bank_name"] or v["bank_code"]:
        hit = np_bank_lookup(v["bank_name"], v["bank_code"])
        if not hit:
            return v, f"'{v['bank_name'] or v['bank_code']}' isn't on the Nepal IPS bank list"
        v["bank_code"], v["bank_name"] = hit[0], hit[1]
    if v["bank_account_number"] and not v["bank_code"]:
        return v, "Account Number given without a Bank Name"
    return v, ""


def plan_changes(rows: list[dict], db) -> dict:
    """Same shape as bank_details_upload.plan_changes. Uploads only ever ADD
    employees: a row whose ledger code, PAN, HEX id or exact name is already
    in nepal_employee_master is EXISTS ("Already in database") and left alone
    — changes go through the Nepal Employees edit form."""
    try:
        existing = db.from_("nepal_employee_master").select("*").execute().data or []
    except Exception:
        existing = []
    index = {k: defaultdict(set) for k in (*_KEYS, "name")}
    by_id = {}
    for e in existing:
        by_id[str(e["id"])] = e
        for k in _KEYS:
            if e.get(k):
                index[k][str(e[k]).strip().upper()].add(str(e["id"]))
        index["name"][_norm_name(e.get("employee_name") or "")].add(str(e["id"]))

    def on_file(v) -> set:
        hits = set(index["name"].get(_norm_name(v["employee_name"]), set())) if v["employee_name"] else set()
        for k in _KEYS:
            if v[k]:
                hits |= index[k].get(v[k].upper(), set())
        return hits

    cleaned = [(r["row_no"], *_normalise(r)) for r in rows]
    cleaned = [(n, v, reason, on_file(v)) for n, v, reason in cleaned]
    seen = defaultdict(int)
    for _, v, reason, hits in cleaned:
        if not reason and not hits:
            for k in _KEYS:
                if v[k]:
                    seen[(k, v[k].upper())] += 1
            if v["employee_name"]:
                seen[("name", _norm_name(v["employee_name"]))] += 1

    items = []
    for row_no, v, reason, hits in cleaned:
        item = {"row_no": row_no, "employee_id": v["employee_ledger_code"] or v["apex_employee_id"] or "",
                "name": v["employee_name"], "fav": "", "changes": [], "warnings": [], "reason": reason}
        if hits:
            names = ", ".join(sorted(by_id[h].get("employee_name") or "?" for h in hits))
            items.append({**item, "action": "EXISTS", "name": names, "reason": "Already in database"})
            continue
        if not reason and not v["employee_name"]:
            reason = "Employee Name is required"
        if not reason:
            dupes = [k for k in _KEYS if v[k] and seen[(k, v[k].upper())] > 1]
            if seen[("name", _norm_name(v["employee_name"]))] > 1:
                dupes.append("name")
            if dupes:
                reason = f"Same {', '.join(FIELD_LABELS.get(d, d) for d in dupes)} appears more than once in this file"
        if reason:
            items.append({**item, "action": "REJECT", "reason": reason})
            continue
        payload = {f: v[f] for f in FIELD_LABELS if v[f]}
        changes = [(FIELD_LABELS[f], v[f]) for f in FIELD_LABELS if v[f]]
        items.append({**item, "action": "NEW", "payload": {"entity": "HNPL", **payload}, "changes": changes})

    order = {"REJECT": 0, "NEW": 1, "EXISTS": 2}
    items.sort(key=lambda i: (order[i["action"]], i["row_no"]))
    return {"items": items, "counts": {a: sum(1 for i in items if i["action"] == a) for a in order}}


def apply_plan(plan: dict, db) -> dict:
    now = datetime.now(timezone.utc).isoformat()
    done = {"NEW": 0}
    for item in plan["items"]:
        if item["action"] == "NEW":
            db.from_("nepal_employee_master").insert({**item["payload"], "source": SOURCE, "imported_at": now}).execute()
            done["NEW"] += 1
    return done


def build_template() -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Nepal Employees"
    ws.append([label for _, label in TEMPLATE_HEADERS])
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F3A5F")
    for i, width in enumerate((30, 14, 16, 20, 14, 12, 34, 30, 22, 16, 14, 14)):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i + 1)].width = width
    # Text format on every ID/number column, so Excel can't drop leading zeros
    # or turn long account numbers into 9.10011E+14.
    for col in "BDEIJK":
        for r in range(2, 1001):
            ws[f"{col}{r}"].number_format = "@"
    ws.freeze_panes = "A2"

    notes = wb.create_sheet("How to fill")
    for line in (
        ["How to fill the Nepal Employees sheet"],
        [],
        ["Employee Name", "Required for someone new."],
        ["Ledger Code / PAN / HEX ID", "If any of these, or the name, is already in APEX, the row is skipped as 'Already in database'."],
        ["Bank Name", "As on the Nepal IPS list, e.g. 'Nabil Bank Ltd.' — common spellings like 'NIMB' work too."],
        ["Account Number", "Type it in this sheet (the column is Text). Don't paste numbers from another sheet — "
                           "Excel drops leading zeros."],
        ["Resign Date", "e.g. 26-Apr-26 or 2026-04-26. Leave blank if still employed."],
        [],
        ["To change someone already in APEX, edit them on the Nepal Employees page."],
    ):
        notes.append(line)
    notes["A1"].font = Font(bold=True, size=13)
    notes.column_dimensions["A"].width = 28
    notes.column_dimensions["B"].width = 100
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
